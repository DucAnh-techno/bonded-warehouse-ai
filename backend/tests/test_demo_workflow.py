import io
import json
import uuid
from datetime import date
import pytest
from openpyxl import load_workbook
from app.models.shipment import Shipment
from app.models.extraction import ExtractedLineItem
from app.models.audit import AuditLog
from app.models.validation import Discrepancy
from app.services.validation.reconciler import reconciliation_engine
from conftest import complete_demo_fixture


@pytest.fixture
def demo(test_db):
    shipment = Shipment(shipment_code='DEMO-'+uuid.uuid4().hex[:8], status='DRAFT')
    test_db.add(shipment); test_db.commit()
    invoice, packing, contract = complete_demo_fixture(test_db, shipment)
    return shipment, invoice, packing, contract


def test_three_document_flow_export_failure_retry_and_duplicate(client, test_db, demo):
    shipment, invoice, packing, contract = demo
    sid = shipment.id
    assert client.get(f'/api/v1/wms/export-excel/{sid}').status_code == 403
    result = client.post(f'/api/v1/validations/reconcile/{sid}').json()
    assert result['error_count'] == 0
    assert client.post(f'/api/v1/wms/approve/{sid}', json={'approved_by':'Demo Teacher'}).status_code == 200
    excel = client.get(f'/api/v1/wms/export-excel/{sid}')
    assert excel.status_code == 200
    workbook = load_workbook(io.BytesIO(excel.content))
    values = [cell.value for row in workbook.active for cell in row]
    assert '0001' in values and 'LEASE-0001' in values
    assert client.post(f'/api/v1/wms/transmit/{sid}', json={'approved_by':'Demo Teacher','simulate_failure':True}).status_code == 503
    assert client.post(f'/api/v1/wms/transmit/{sid}', json={'approved_by':'Demo Teacher'}).status_code == 200
    audit_count = test_db.query(AuditLog).filter_by(shipment_id=sid).count()
    client.get(f'/api/v1/validations/reconcile/{sid}')
    test_db.refresh(shipment)
    assert shipment.status == 'TRANSMITTED'
    assert test_db.query(AuditLog).filter_by(shipment_id=sid).count() == audit_count
    assert client.post(f'/api/v1/wms/transmit/{sid}', json={'approved_by':'Demo Teacher'}).status_code == 200


@pytest.mark.parametrize('change', ['start_after_entry','wrong_owner','missing_number','missing_date'])
def test_contract_errors_block_approval_and_export(client, test_db, demo, change):
    shipment, invoice, packing, contract = demo
    hdr = json.loads(contract.extraction.extracted_header_json)
    if change == 'start_after_entry': hdr['valid_from'] = '2026-12-01'  # After entry_date (2026-01-05)
    if change == 'wrong_owner': hdr['tenant_name'] = 'OTHER OWNER'
    if change == 'missing_number': hdr['document_number'] = None
    if change == 'missing_date': shipment.entry_date = None
    contract.extraction.extracted_header_json = json.dumps(hdr)
    test_db.commit()
    response = client.post(f'/api/v1/wms/approve/{shipment.id}',json={'approved_by':'Teacher'})
    assert response.status_code == 400
    assert client.get(f'/api/v1/wms/export-excel/{shipment.id}').status_code == 403


def test_same_sku_multiple_rows_and_different_po_are_not_overwritten(client, test_db, demo):
    shipment, invoice, packing, contract = demo
    invoice.extraction.line_items[0].po_reference = 'PO-1'
    packing.extraction.line_items[0].po_reference = 'PO-1'
    for doc in (invoice, packing):
        for po,qty in [('PO-1',3),('PO-2',7)]:
            doc.extraction.line_items.append(ExtractedLineItem(sku='0001',description_vi='Demo',quantity=qty,uom='PCS',po_reference=po,unit_price=3,total_amount=qty*3))
    hdr=json.loads(invoice.extraction.extracted_header_json);hdr['total_invoice_value']=36
    invoice.extraction.extracted_header_json=json.dumps(hdr)
    test_db.commit()
    result=client.post(f'/api/v1/validations/reconcile/{shipment.id}').json()
    assert result['error_count'] == 0
    assert {(r['po_reference'],r['qty_invoice']) for r in result['sku_matrix']} == {('PO-1',5),('PO-2',7)}
    packing.extraction.line_items[-1].quantity = 8
    test_db.commit()
    result=client.post(f'/api/v1/validations/reconcile/{shipment.id}').json()
    assert result['error_count'] > 0


def test_header_correction_audit_and_reapproval(client, test_db, demo):
    shipment, invoice, packing, contract = demo
    client.post(f'/api/v1/wms/approve/{shipment.id}',json={'approved_by':'Teacher'})
    response=client.patch(f'/api/v1/extractions/document/{contract.id}/header',json={'values':{'valid_from':'2030-01-01'},'reason':'Test invalid start lease'})
    assert response.status_code == 200
    test_db.refresh(shipment)
    assert shipment.approval_status == 'PENDING'
    assert shipment.status == 'DISCREPANT'
    assert test_db.query(AuditLog).filter_by(target_id=contract.extraction.id,field_name='valid_from').count() == 1
    assert client.patch(f'/api/v1/shipments/{shipment.id}',json={'status':'VALIDATED','approval_status':'APPROVED'}).status_code == 422


def test_uploaded_but_not_extracted_is_not_valid(client,test_db,demo):
    shipment,invoice,packing,contract=demo
    contract.extraction.status='FAILED';test_db.commit()
    assert client.post(f'/api/v1/wms/approve/{shipment.id}',json={'approved_by':'Teacher'}).status_code == 400


def test_batch_failure_is_saved_and_blocks_old_results(client, test_db, demo, monkeypatch, tmp_path):
    from app.api.v1 import extractions
    shipment, invoice, packing, contract = demo
    sample = tmp_path / 'sample.pdf'
    sample.write_bytes(b'%PDF-test')
    monkeypatch.setattr(extractions.storage_service, 'get_absolute_path', lambda path: sample)
    async def fail(**kwargs):
        raise RuntimeError('503 provider unavailable')
    monkeypatch.setattr(extractions.gemini_extractor, 'batch_extract_and_reconcile_shipment', fail)
    assert client.post(f'/api/v1/wms/approve/{shipment.id}',json={'approved_by':'Teacher'}).status_code == 200
    result = client.post(f'/api/v1/extractions/shipment/{shipment.id}/batch-reconcile')
    assert result.status_code == 500
    for doc in (invoice, packing, contract):
        saved = client.get(f'/api/v1/extractions/document/{doc.id}').json()
        assert saved['status'] == 'FAILED' and saved['error_message']
    assert client.get(f'/api/v1/wms/export-excel/{shipment.id}').status_code == 403


def test_empty_shipment_cannot_be_approved(client):
    sid=client.post('/api/v1/shipments/',json={}).json()['id']
    assert client.post(f'/api/v1/wms/approve/{sid}',json={'approved_by':'Teacher'}).status_code == 400
    result=client.post(f'/api/v1/validations/shipment/{sid}/approve',json={'approved_by':'Teacher'}).json()
    assert result['success'] is False


def test_retired_endpoints_are_absent(client):
    paths=client.get('/openapi.json').json()['paths']
    assert not any('benchmarks' in p or 'ai-explain' in p for p in paths)


def test_missing_quantity_remains_null_and_blocks_wms(client, test_db, demo):
    shipment, invoice, packing, contract = demo
    item = invoice.extraction.line_items[0]
    response = client.patch(f'/api/v1/extractions/items/{item.id}', json={'quantity':None, 'reason':'Missing on document'})
    assert response.status_code == 200
    assert response.json()['quantity'] is None
    assert client.post(f'/api/v1/wms/approve/{shipment.id}',json={'approved_by':'Teacher'}).status_code == 400


def test_demo_mode_rejects_other_files_and_enforces_ai_budget(client, monkeypatch):
    from app.core.config import settings
    from app.core import demo_limits
    from pathlib import Path
    monkeypatch.setattr(settings, 'DEMO_ONLY', True)
    sid=client.post('/api/v1/shipments/',json={}).json()['id']
    response=client.post('/api/v1/documents/upload',data={'shipment_id':sid,'document_type':'COMMERCIAL_INVOICE'},files={'file':('private.pdf',b'not-a-sample','application/pdf')})
    assert response.status_code == 400
    sample=settings.DEMO_SAMPLES_DIR/'Invoice-demo.pdf'
    response=client.post('/api/v1/documents/upload',data={'shipment_id':sid,'document_type':'COMMERCIAL_INVOICE'},files={'file':(sample.name,sample.read_bytes(),'application/pdf')})
    assert response.status_code == 201
    monkeypatch.setattr(settings,'DEMO_AI_REQUESTS_PER_HOUR',0)
    assert client.post('/api/v1/extractions/extract/'+response.json()['id']).status_code == 429


def test_multiple_invoice_exports_and_text_cells(client,test_db,demo):
    from app.models.document import Document
    from app.models.extraction import DocumentExtraction
    from app.services.export.payload_builder import payload_builder
    from app.services.export.excel_mapper import wms_excel_mapper
    shipment,invoice,packing,contract=demo
    other=Document(shipment_id=shipment.id,document_type='COMMERCIAL_INVOICE',original_filename='second.pdf',storage_path=uuid.uuid4().hex+'.pdf',mime_type='application/pdf',sha256_hash='test',file_size_bytes=1)
    test_db.add(other);test_db.flush()
    other.extraction=DocumentExtraction(document_id=other.id,status='SUCCESS',model_name='test',raw_response_json='{}',extracted_header_json=json.dumps({'consignee':'DEMO TENANT','total_invoice_value':3}))
    other.extraction.line_items.append(ExtractedLineItem(sku='=1+1',description_vi='Demo',quantity=1,uom='PCS',unit_price=3,total_amount=3))
    packing.extraction.line_items.append(ExtractedLineItem(sku='=1+1',description_vi='Demo',quantity=1,uom='PCS'))
    test_db.commit()
    assert client.post(f'/api/v1/wms/approve/{shipment.id}',json={'approved_by':'Teacher'}).status_code==200
    assert len(payload_builder.build_inbound_receipt_payload(test_db,shipment).line_items)==2
    book=load_workbook(io.BytesIO(wms_excel_mapper.generate_wms_excel(test_db,shipment.id)))
    cell=next(c for row in book.active for c in row if c.value=='=1+1')
    assert cell.data_type=='s'


def test_repeated_validation_does_not_duplicate_open_findings(client,test_db,demo):
    shipment,invoice,packing,contract=demo
    packing.extraction.line_items[0].quantity=99;test_db.commit()
    for _ in range(3):client.post(f'/api/v1/validations/reconcile/{shipment.id}')
    assert test_db.query(Discrepancy).filter_by(shipment_id=shipment.id,rule_code='VAL-QTY-001',status='OPEN').count()==1
